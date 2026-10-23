"""Resume parsing (PDF/DOCX/TXT), drafting from profile, and export (PDF/DOCX)."""
import io

from .data import split_list


def extract_text(uploaded_name: str, data: bytes) -> str:
    name = uploaded_name.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".docx"):
        import docx
        return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    return data.decode("utf-8", errors="ignore")


def draft_from_profile(p: dict, target_role: str = "") -> str:
    """Deterministic resume using ONLY the user's own details."""
    lines = [p.get("name") or "Your Name"]
    contact = " | ".join(x for x in [p.get("email"), p.get("phone"), p.get("linkedin")] if x)
    if contact: lines.append(contact)
    lines.append("")
    role = target_role or ", ".join(split_list(p.get("target_roles"))[:1])
    lines += ["SUMMARY", f"{p.get('education','')} {('in ' + p['specialization']) if p.get('specialization') else ''} graduate"
              f"{' seeking a ' + role + ' role' if role else ''}. {p.get('goals','')}".strip(), ""]
    if p.get("skills"): lines += ["SKILLS", ", ".join(split_list(p["skills"])), ""]
    for header, key in [("EXPERIENCE", "experience"), ("PROJECTS", "projects"), ("CERTIFICATIONS", "certifications")]:
        if p.get(key):
            lines.append(header)
            lines += [f"- {x.strip()}" for x in p[key].split("\n") if x.strip()]
            lines.append("")
    edu = ", ".join(x for x in [p.get("education"), p.get("specialization"), p.get("institution"), str(p.get("grad_year") or "")] if x)
    if edu: lines += ["EDUCATION", edu]
    return "\n".join(lines).strip()


def to_docx(text: str) -> bytes:
    import docx
    d = docx.Document()
    for i, line in enumerate(text.split("\n")):
        if i == 0:
            d.add_heading(line, level=0)
        elif line.isupper() and line.strip():
            d.add_heading(line.title(), level=2)
        elif line.startswith("- "):
            d.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pdf(text: str) -> bytes:
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(True, 15)
    for i, line in enumerate(text.split("\n")):
        safe = line.encode("latin-1", "replace").decode("latin-1")
        if i == 0:
            pdf.set_font("Helvetica", "B", 18); pdf.multi_cell(0, 10, safe, new_x="LMARGIN", new_y="NEXT")
        elif safe.isupper() and safe.strip():
            pdf.ln(2); pdf.set_font("Helvetica", "B", 12); pdf.multi_cell(0, 7, safe, new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.set_font("Helvetica", "", 10); pdf.multi_cell(0, 5.5, safe, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
